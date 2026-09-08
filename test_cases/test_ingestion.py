import tempfile
import unittest
from pathlib import Path
import docx
import fitz
from pptx import Presentation
from pptx.util import Inches

from app.database.repository import DatabaseManager
from app.embeddings.mock import MockEmbeddingModel
from app.ingestion.docx import DOCXExtractor
from app.ingestion.pdf import PDFExtractor
from app.ingestion.pipeline import IngestionPipeline
from app.ingestion.pptx import PPTXExtractor
from app.vector_store.chroma import ChromaVectorStore


class TestIngestion(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory(ignore_cleanup_errors=True)
        self.base_path = Path(self.temp_dir.name)

        # 1. Create a 2-page text PDF
        self.pdf_path = self.base_path / "sample_manual.pdf"
        pdf_doc = fitz.open()
        p1 = pdf_doc.new_page()
        p1.insert_text((50, 72), "Chapter 1: SOAR Sovereign System Architecture and Overview.")
        p2 = pdf_doc.new_page()
        p2.insert_text((50, 72), "Chapter 2: Security compliance, air-gapped guarantees, and audit.")
        pdf_doc.save(self.pdf_path)
        pdf_doc.close()

        # 2. Create a scanned (blank / textless) PDF
        self.scanned_pdf_path = self.base_path / "scanned_invoice.pdf"
        scanned_doc = fitz.open()
        scanned_doc.new_page()  # empty raster-like page with zero text
        scanned_doc.save(self.scanned_pdf_path)
        scanned_doc.close()

        # 3. Create a DOCX document
        self.docx_path = self.base_path / "policy.docx"
        docx_doc = docx.Document()
        docx_doc.add_heading("SOAR Operational Policy", level=1)
        docx_doc.add_paragraph("All models must run strictly on-premise without cloud transmission.")
        table = docx_doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Role"
        table.cell(0, 1).text = "Access Level"
        table.cell(1, 0).text = "Admin"
        table.cell(1, 1).text = "Full"
        docx_doc.save(self.docx_path)

        # 4. Create a PPTX presentation
        self.pptx_path = self.base_path / "briefing.pptx"
        prs = Presentation()
        slide1 = prs.slides.add_slide(prs.slide_layouts[0])
        slide1.shapes.title.text = "SOAR SIH 2026 Project"
        slide1.placeholders[1].text = "Sovereign On-premise Agentic Reasoning"
        slide2 = prs.slides.add_slide(prs.slide_layouts[1])
        slide2.shapes.title.text = "Data Layer Capabilities"
        slide2.placeholders[1].text = "SQLite, ChromaDB, BGE-M3 Multimodal Pipeline"
        prs.save(self.pptx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pdf_extractor_and_page_metadata(self):
        extractor = PDFExtractor()
        res = extractor.extract(self.pdf_path)

        self.assertEqual(res.status, "success")
        self.assertEqual(res.file_type, "pdf")
        self.assertEqual(len(res.elements), 2)
        self.assertEqual(res.elements[0].metadata["page_number"], 1)
        self.assertIn("SOAR Sovereign System Architecture", res.elements[0].text)
        self.assertEqual(res.elements[1].metadata["page_number"], 2)
        self.assertIn("Security compliance", res.elements[1].text)
        self.assertFalse(res.is_scanned)

    def test_scanned_pdf_detection(self):
        extractor = PDFExtractor()
        res = extractor.extract(self.scanned_pdf_path)

        self.assertEqual(res.status, "scanned_needs_ocr")
        self.assertTrue(res.is_scanned)
        self.assertTrue(res.needs_ocr)
        self.assertEqual(len(res.elements), 1)
        self.assertIn("[Scanned PDF Document", res.elements[0].text)
        self.assertTrue(res.elements[0].metadata["is_scanned"])

    def test_docx_extractor(self):
        extractor = DOCXExtractor()
        res = extractor.extract(self.docx_path)

        self.assertEqual(res.status, "success")
        self.assertEqual(res.file_type, "docx")
        self.assertTrue(len(res.elements) >= 2)
        # Check paragraph & table extraction
        texts = [e.text for e in res.elements]
        self.assertTrue(any("SOAR Operational Policy" in t for t in texts))
        self.assertTrue(any("All models must run strictly" in t for t in texts))
        self.assertTrue(any("Role | Access Level" in t for t in texts))

    def test_pptx_extractor(self):
        extractor = PPTXExtractor()
        res = extractor.extract(self.pptx_path)

        self.assertEqual(res.status, "success")
        self.assertEqual(res.file_type, "pptx")
        self.assertEqual(len(res.elements), 2)
        self.assertEqual(res.elements[0].metadata["slide_number"], 1)
        self.assertIn("SOAR SIH 2026 Project", res.elements[0].text)
        self.assertEqual(res.elements[1].metadata["slide_number"], 2)
        self.assertIn("Data Layer Capabilities", res.elements[1].text)

    def test_end_to_end_pipeline_and_duplicate_detection(self):
        db_path = self.base_path / "test_pipeline.db"
        chroma_dir = self.base_path / "chroma"
        db = DatabaseManager(db_path=db_path)
        embedding_model = MockEmbeddingModel(dimension=256)
        vector_store = ChromaVectorStore(
            collection_name="test_pipeline_kb",
            persist_directory=chroma_dir,
            embedding_model=embedding_model,
        )

        pipeline = IngestionPipeline(
            db_manager=db,
            vector_store=vector_store,
            embedding_model=embedding_model,
        )

        # 1. Ingest PDF
        pdf_res = pipeline.ingest_file(self.pdf_path)
        self.assertEqual(pdf_res["status"], "success")
        self.assertEqual(pdf_res["chunks_indexed"], 2)

        # 2. Ingest DOCX
        docx_res = pipeline.ingest_file(self.docx_path)
        self.assertEqual(docx_res["status"], "success")
        self.assertTrue(docx_res["chunks_indexed"] >= 2)

        # 3. Ingest PPTX
        pptx_res = pipeline.ingest_file(self.pptx_path)
        self.assertEqual(pptx_res["status"], "success")
        self.assertEqual(pptx_res["chunks_indexed"], 2)

        # 4. Ingest Scanned PDF
        scanned_res = pipeline.ingest_file(self.scanned_pdf_path)
        self.assertEqual(scanned_res["status"], "scanned_needs_ocr")
        self.assertTrue(scanned_res["is_scanned"])

        # Check total documents in SQLite
        docs = db.list_documents()
        self.assertEqual(len(docs), 4)

        # 5. Test Duplicate Detection (re-ingesting pdf_path without force)
        dup_res = pipeline.ingest_file(self.pdf_path, force_reingest=False)
        self.assertEqual(dup_res["status"], "duplicate")
        self.assertEqual(dup_res["chunks_indexed"], 0)
        self.assertEqual(dup_res["document_id"], pdf_res["document_id"])
        # SQLite document count remains 4
        self.assertEqual(len(db.list_documents()), 4)

        # 6. Verify vector similarity search across indexed documents
        search_results = vector_store.similarity_search("Chapter 1: SOAR Sovereign System Architecture and Overview.", k=2)
        self.assertTrue(len(search_results) > 0)
        top_meta = search_results[0]["metadata"]
        self.assertEqual(top_meta["filename"], "sample_manual.pdf")
        self.assertEqual(top_meta["page_number"], 1)

        # 7. Check audit logs
        audit_logs = db.get_audit_logs()
        self.assertTrue(len(audit_logs) >= 4)

    def test_directory_ingestion(self):
        db_path = self.base_path / "test_dir.db"
        chroma_dir = self.base_path / "chroma_dir"
        db = DatabaseManager(db_path=db_path)
        embedding_model = MockEmbeddingModel(dimension=256)
        vector_store = ChromaVectorStore(
            collection_name="test_dir_kb",
            persist_directory=chroma_dir,
            embedding_model=embedding_model,
        )

        pipeline = IngestionPipeline(
            db_manager=db,
            vector_store=vector_store,
            embedding_model=embedding_model,
        )

        batch_results = pipeline.ingest_directory(self.base_path)
        self.assertTrue(len(batch_results) >= 4)
        self.assertTrue(all(r["status"] in ("success", "scanned_needs_ocr") for r in batch_results))


if __name__ == "__main__":
    unittest.main()
