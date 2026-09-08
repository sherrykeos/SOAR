import json
import os
import tempfile
from pathlib import Path

import docx

from app.models.adapter import ModelCapability
from app.models.manager import ModelManager
from app.models.mock import MockModel
from app.models.registry import ModelRegistry
from app.models.router import ModelRouter
from app.orchestrator.action import ToolAction
from app.orchestrator.agent import Agent
from app.orchestrator.executor import Executor
from app.orchestrator.planner import Planner
from app.tools.docx_creator import DOCXCreatorTool
from app.tools.registry import ToolRegistry


print("[TEST] Starting DOCXCreatorTool test suite...\n")

with tempfile.TemporaryDirectory() as temp_dir:
    temp_dir_path = temp_dir.replace("\\", "/")
    output_docx = f"{temp_dir_path}/test_approval_note.docx"
    multi_p_docx = f"{temp_dir_path}/multi_paragraph.docx"

    # 1. Registration & Retrieval
    print("1 & 2. Testing ToolRegistry registration and retrieval...")
    registry = ToolRegistry()
    registry.register(DOCXCreatorTool())

    tool = registry.get("docx_creator")
    assert tool.name == "docx_creator"
    assert "DOCX" in tool.description
    print(f"Successfully retrieved tool '{tool.name}'.")

    # 2. DOCX creation
    print("\n3. Testing DOCX creation with title and content...")
    result = tool.execute(
        output_path=output_docx,
        title="Industrial Plant Approval Note",
        content="The inspection confirmed all systems are within safety margins.",
    )
    print(f"Execution result: {result}")
    assert "Successfully created" in result
    assert os.path.exists(output_docx), "DOCX file was not created on disk!"
    print("DOCX file exists on disk.")

    # 3. Verify DOCX structure & contents using python-docx
    print("\n4, 5, 6. Verifying created DOCX file structure and content...")
    doc = docx.Document(output_docx)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    all_text = [p.text for p in doc.paragraphs]

    print(f"Document paragraphs: {all_text}")
    assert any("Industrial Plant Approval Note" in text for text in all_text)
    assert any("The inspection confirmed all systems" in text for text in all_text)
    print("DOCX title and body content verified successfully.")

    # 4. Multi-paragraph support
    print("\n7. Testing multi-paragraph content support...")
    multi_content = (
        "Paragraph 1: Boiler pressure inspection complete.\n\n"
        "Paragraph 2: Turbine vibration levels normal.\n\n"
        "Paragraph 3: Approved for operation."
    )
    multi_res = tool.execute(
        output_path=multi_p_docx,
        title="Comprehensive Equipment Report",
        content=multi_content,
    )
    assert "Successfully created" in multi_res
    doc_multi = docx.Document(multi_p_docx)
    paragraphs = [p.text for p in doc_multi.paragraphs if p.text.strip()]
    assert len(paragraphs) >= 4  # 1 title heading + 3 body paragraphs
    assert any("Paragraph 1" in p for p in paragraphs)
    assert any("Paragraph 2" in p for p in paragraphs)
    assert any("Paragraph 3" in p for p in paragraphs)
    print("Multi-paragraph DOCX generation verified successfully.")

    # 5. Missing parameter & Error handling
    print("\n8. Testing missing parameter handling...")
    missing_res = tool.execute(title="No output path")
    print(f"Missing parameter result: {missing_res}")
    assert "Missing required parameter" in missing_res

    # 6. Executor Integration Test
    print("\n9. Testing execution through Executor...")
    executor = Executor(tool_registry=registry)
    exec_docx = f"{temp_dir_path}/exec_test.docx"
    action = ToolAction(
        tool="docx_creator",
        arguments={
            "output_path": exec_docx,
            "title": "Executor Note",
            "content": "Generated through Executor.",
        },
    )
    exec_result = executor.execute_action(action)
    print(f"Executor result: {exec_result}")
    assert exec_result["status"] == "completed"
    assert os.path.exists(exec_docx)
    print("Executor execution verified.")

    # 7. Agent Loop Integration Test
    print("\n10. Testing Agent Loop integration with DOCXCreatorTool...")
    agent_docx = f"{temp_dir_path}/agent_approval.docx"

    class DOCXSequenceResponder:
        def __init__(self, docx_path: str):
            self.count = 0
            self.docx_path = docx_path

        def __call__(self, prompt: str) -> str:
            self.count += 1
            if self.count == 1:
                return json.dumps({
                    "status": "continue",
                    "thought": "I will create an approval note in DOCX format.",
                    "steps": [
                        {
                            "tool": "docx_creator",
                            "arguments": {
                                "output_path": self.docx_path,
                                "title": "Final Approval Note",
                                "content": "All compliance criteria met.",
                            }
                        }
                    ]
                })
            else:
                return json.dumps({
                    "status": "complete",
                    "thought": "Approval note DOCX has been created. Task complete.",
                    "steps": []
                })

    mock_model = MockModel(
        model_id="mock-docx-agent",
        capabilities={ModelCapability.GENERAL, ModelCapability.REASONING},
        priority=1,
        response_fn=DOCXSequenceResponder(agent_docx),
    )

    model_registry = ModelRegistry()
    model_registry.register(mock_model)
    model_router = ModelRouter(model_registry, default_model_id="mock-docx-agent")
    model_manager = ModelManager(model_registry=model_registry, model_router=model_router)

    planner = Planner(model_manager=model_manager, tool_registry=registry)
    agent = Agent(planner=planner, executor=executor, max_iterations=5)

    agent_state = agent.run("Create approval note DOCX.")
    print(f"Agent Final Status: {agent_state.status}")
    print(f"Agent Observations: {agent_state.observations}")

    assert agent_state.status == "completed"
    assert agent_state.iterations == 2
    assert os.path.exists(agent_docx)
    assert agent_state.observations[0]["tool"] == "docx_creator"
    print("Agent DOCX integration verified successfully!")

print("\n[TEST] All DOCXCreatorTool tests completed successfully!")
