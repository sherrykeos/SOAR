import hashlib
import uuid
from pathlib import Path
from typing import List

from .base import BaseExtractor, DocumentIngestionResult, ExtractedElement


class PPTXExtractor(BaseExtractor):
    """
    Extracts text, slide titles, shape content, and speaker notes from PPTX presentations.
    Preserves slide-level indexing and structural metadata.
    """

    @property
    def supported_extensions(self) -> List[str]:
        return [".pptx"]

    def extract(self, file_path: str | Path) -> DocumentIngestionResult:
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"PPTX file not found: {path}")

        file_bytes = path.read_bytes()
        content_hash = hashlib.sha256(file_bytes).hexdigest()
        file_size = len(file_bytes)
        filename = path.name
        doc_id = str(uuid.uuid4())

        try:
            from pptx import Presentation
        except ImportError as e:
            raise ImportError(
                "python-pptx is required for PPTX extraction. Install it via `pip install python-pptx`."
            ) from e

        try:
            prs = Presentation(path)
            elements: List[ExtractedElement] = []
            total_slides = len(prs.slides)

            for slide_idx, slide in enumerate(prs.slides):
                slide_num = slide_idx + 1
                slide_texts: List[str] = []

                # Extract text from shapes
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            text = (paragraph.text or "").strip()
                            if text:
                                slide_texts.append(text)
                    elif shape.has_table:
                        for row in shape.table.rows:
                            row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                            if row_vals:
                                slide_texts.append(" | ".join(row_vals))

                # Extract speaker notes if present
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    notes_text = (slide.notes_slide.notes_text_frame.text or "").strip()
                    if notes_text:
                        slide_texts.append(f"[Speaker Notes]: {notes_text}")

                if slide_texts:
                    combined_slide_text = "\n".join(slide_texts)
                    elements.append(
                        ExtractedElement(
                            text=combined_slide_text,
                            metadata={
                                "filename": filename,
                                "source_path": str(path),
                                "file_type": "pptx",
                                "slide_number": slide_num,
                                "total_slides": total_slides,
                                "content_type": "slide",
                            },
                        )
                    )

            return DocumentIngestionResult(
                document_id=doc_id,
                filename=filename,
                source_path=str(path),
                file_type="pptx",
                file_size=file_size,
                content_hash=content_hash,
                elements=elements,
                is_scanned=False,
                needs_ocr=False,
                status="success",
            )

        except Exception as e:
            return DocumentIngestionResult(
                document_id=doc_id,
                filename=filename,
                source_path=str(path),
                file_type="pptx",
                file_size=file_size,
                content_hash=content_hash,
                elements=[],
                status="error",
                error=f"Error extracting PPTX '{filename}': {str(e)}",
            )
