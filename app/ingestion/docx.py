import hashlib
import uuid
from pathlib import Path
from typing import List

from .base import BaseExtractor, DocumentIngestionResult, ExtractedElement


class DOCXExtractor(BaseExtractor):
    """
    Extracts text, paragraphs, and tables from DOCX documents using python-docx.
    Preserves structural section metadata.
    """

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    def extract(self, file_path: str | Path) -> DocumentIngestionResult:
        path = Path(file_path).resolve()
        if not path.exists():
            raise FileNotFoundError(f"DOCX file not found: {path}")

        file_bytes = path.read_bytes()
        content_hash = hashlib.sha256(file_bytes).hexdigest()
        file_size = len(file_bytes)
        filename = path.name
        doc_id = str(uuid.uuid4())

        try:
            import docx
        except ImportError as e:
            raise ImportError(
                "python-docx is required for DOCX extraction. Install it via `pip install python-docx`."
            ) from e

        try:
            doc = docx.Document(path)
            elements: List[ExtractedElement] = []

            # 1. Extract paragraphs
            for p_idx, para in enumerate(doc.paragraphs):
                p_text = (para.text or "").strip()
                if p_text:
                    style_name = para.style.name if para.style else "Normal"
                    elements.append(
                        ExtractedElement(
                            text=p_text,
                            metadata={
                                "filename": filename,
                                "source_path": str(path),
                                "file_type": "docx",
                                "paragraph_index": p_idx + 1,
                                "style": style_name,
                                "content_type": "paragraph",
                            },
                        )
                    )

            # 2. Extract tables
            for t_idx, table in enumerate(doc.tables):
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells]
                    if any(row_cells):
                        table_rows.append(" | ".join(row_cells))

                if table_rows:
                    table_text = "\n".join(table_rows)
                    elements.append(
                        ExtractedElement(
                            text=table_text,
                            metadata={
                                "filename": filename,
                                "source_path": str(path),
                                "file_type": "docx",
                                "table_index": t_idx + 1,
                                "content_type": "table",
                            },
                        )
                    )

            if not elements:
                # Handle empty docx
                return DocumentIngestionResult(
                    document_id=doc_id,
                    filename=filename,
                    source_path=str(path),
                    file_type="docx",
                    file_size=file_size,
                    content_hash=content_hash,
                    elements=[],
                    is_scanned=False,
                    needs_ocr=False,
                    status="success",
                )

            return DocumentIngestionResult(
                document_id=doc_id,
                filename=filename,
                source_path=str(path),
                file_type="docx",
                file_size=file_size,
                content_hash=content_hash,
                elements=elements,
                is_scanned=False,
                needs_ocr=False,
                status="success",
            )

        except Exception as e:
            return DocumentIngestionResult(
                document_id=doc_id,
                filename=filename,
                source_path=str(path),
                file_type="docx",
                file_size=file_size,
                content_hash=content_hash,
                elements=[],
                status="error",
                error=f"Error extracting DOCX '{filename}': {str(e)}",
            )
